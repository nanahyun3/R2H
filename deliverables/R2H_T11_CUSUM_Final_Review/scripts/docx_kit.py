"""python-docx 보조 — 한글 글꼴(eastAsia), OMML 수식, 표·그림 캡션, 머리말·꼬리말, 목차 필드.

수식: LaTeX → MathML (latex2mathml) → OMML (Office 의 MML2OMML.XSL) — Word 에서 편집 가능한 수식 객체로 들어간다.
"""

from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree

BODY_FONT = "맑은 고딕"
BULLET_SIZE: list[float | None] = [None]
LATIN_FONT = "Malgun Gothic"
XSL = Path(os.environ.get("MML2OMML_XSL", "/mnt/c/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL"))
_XSLT = etree.XSLT(etree.parse(str(XSL)))


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None, italic=None) -> None:
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _style_font(style, size, bold=False, color=None, space_before=0, space_after=4, line=1.3):
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = style.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line


def new_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1.1)
    st = doc.styles
    _style_font(st["Normal"], 10, space_after=5, line=1.35)
    _style_font(st["Heading 1"], 15, True, "1F3864", space_before=6, space_after=8, line=1.2)
    _style_font(st["Heading 2"], 12, True, "1F3864", space_before=10, space_after=5, line=1.2)
    _style_font(st["Heading 3"], 10.5, True, "2E2E2E", space_before=8, space_after=3, line=1.2)
    _style_font(st["Caption"], 8.5, True, "1F1F1F", space_before=3, space_after=2, line=1.15)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        st[name].paragraph_format.keep_with_next = True
    _style_font(st["List Bullet"], 10, space_after=2, line=1.3)
    _style_font(st["Title"], 22, True, "1F3864", space_after=6, line=1.15)
    return doc


def para(doc, text: str = "", *, size=None, bold=False, align=None, italic=False, color=None, style=None, space_after=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_rich(p, text, size=size, bold=bold, italic=italic, color=color)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


# 본문 속 기호 → 인라인 OMML (결합 악센트 ̂ 는 한글 글꼴에서 깨지므로 수식 객체로 넣는다). 긴 것부터 맞춘다.
MATH_TOKENS: dict[str, str] = {
    "σ̂²_{t−1}": r"\hat{\sigma}_{t-1}^2", "σ̃²_{t−1}": r"\tilde{\sigma}_{t-1}^2", "y²_{t−1}": r"y_{t-1}^2", "y²_{t−4}": r"y_{t-4}^2",
    "σ̂²_t": r"\hat{\sigma}_t^2", "σ̃²_t": r"\tilde{\sigma}_t^2", "ε̂_t": r"\hat{\varepsilon}_t", "σ̂_t": r"\hat{\sigma}_t",
    "ε²_t": r"\varepsilon_t^2", "y²_t": r"y_t^2", "T_t(w)": r"T_t(w)", "T_t(60)": r"T_t(60)",
    "ε̂²": r"\hat{\varepsilon}^2", "σ̂²": r"\hat{\sigma}^2", "τ̂²": r"\hat{\tau}^2", "σ̂": r"\hat{\sigma}", "τ̂": r"\hat{\tau}",
    "k̂": r"\hat{k}", "T_n": r"T_n", "ĝ": r"\hat{g}",
}
_MATH_RE = re.compile("(" + "|".join(re.escape(k) for k in sorted(MATH_TOKENS, key=len, reverse=True)) + ")")


def add_rich(p, text: str, *, size=None, bold=False, italic=False, color=None) -> None:
    """**굵게** 와 수식 기호(MATH_TOKENS)를 해석한다."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        for seg in _MATH_RE.split(part):
            if not seg:
                continue
            if seg in MATH_TOKENS:
                latex = MATH_TOKENS[seg]
                m = omml(latex)
                if size:  # 수식 글자 크기를 주변 글과 맞춘다
                    for rr in m.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}r"):
                        rpr = OxmlElement("w:rPr")
                        sz = OxmlElement("w:sz")
                        sz.set(qn("w:val"), str(int(size * 2)))
                        rpr.append(sz)
                        rr.insert(0, rpr)
                p._p.append(m)
                continue
            r = p.add_run(seg)
            set_run_font(r, size=size, bold=bold or (i % 2 == 1), italic=italic, color=color)


def bullets(doc, items: list[str], size=None) -> None:
    size = size or BULLET_SIZE[0]
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        add_rich(p, it, size=size)


def omml(latex: str):
    mml = etree.fromstring(latex_to_mathml(latex).encode())
    return _XSLT(mml).getroot()


def equation(doc, latex: str, number: int | None) -> None:
    """가운데 정렬 수식 + 오른쪽 수식 번호 (탭 정렬)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    width = Cm(21.0 - 4.6)
    pf.tab_stops.add_tab_stop(int(width / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    pf.space_before, pf.space_after = Pt(4), Pt(6)
    pf.keep_together = True
    r = p.add_run("\t")
    set_run_font(r)
    p._p.append(omml(latex))
    r = p.add_run(f"\t({number})" if number else "\t")
    set_run_font(r, size=9.5)


def inline_math(p, latex: str) -> None:
    p._p.append(omml(latex))


def _shade(cell, fill: str) -> None:
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _cant_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def caption(doc, text: str, keep_next: bool = False):
    p = doc.add_paragraph(style="Caption")
    add_rich(p, text, size=8.8, bold=True)
    p.paragraph_format.keep_with_next = keep_next
    return p


def note(doc, text: str, size: float = 7.8):
    p = doc.add_paragraph()
    add_rich(p, text, size=size, color="404040")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    return p


def table(doc, cap: str, header: list[str], rows: list[list], *, widths: list[float] | None = None,
          size: float = 8.0, src: str | None = None, align: list[str] | None = None) -> None:
    caption(doc, cap, keep_next=True)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        add_rich(c.paragraphs[0], str(h), size=size, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(c, "DCE4F0")
    _repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pp = cells[i].paragraphs[0]
            add_rich(pp, "" if v is None else str(v), size=size)
            a = (align or ["l"] * len(header))[i]
            pp.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}.get(a, WD_ALIGN_PARAGRAPH.LEFT)
            pp.paragraph_format.space_after = Pt(0)
            pp.paragraph_format.line_spacing = 1.1
    for r in t.rows:
        _cant_split(r)
        for c in r.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
    for r in t.rows[:-1]:
        for c in r.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = True
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    if src:
        note(doc, src)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, path: Path, cap: str, src: str, width_cm: float = 16.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    caption(doc, cap, keep_next=True)
    note(doc, src)


def page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _field(run, instr: str, placeholder: str = "") -> None:
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def toc(doc) -> None:
    p = doc.add_paragraph()
    r = p.add_run()
    set_run_font(r, size=10)
    _field(r, 'TOC \\o "1-2" \\h \\z \\u', "목차는 Word 에서 필드 갱신 시 생성됩니다.")


def header_footer(doc, header_text: str, footer_left: str) -> None:
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = ""
    add_rich(hp, header_text, size=7.5, color="555555")
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _bottom_border(hp)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), WD_TAB_ALIGNMENT.RIGHT)
    add_rich(fp, footer_left + "\t", size=7.5, color="555555")
    r = fp.add_run()
    set_run_font(r, size=8)
    _field(r, "PAGE", "1")
    r = fp.add_run(" / ")
    set_run_font(r, size=8)
    r = fp.add_run()
    set_run_font(r, size=8)
    _field(r, "NUMPAGES", "1")


def _bottom_border(p) -> None:
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "999999")):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def set_core(doc, title: str, subject: str, author: str, keywords: str) -> None:
    cp = doc.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = title, subject, author, keywords
    cp.language = "ko-KR"
