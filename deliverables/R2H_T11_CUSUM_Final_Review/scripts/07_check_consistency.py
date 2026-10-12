"""최종 일치성 검수 — claims.json ↔ DOCX ↔ PDF.

1. 본문에서 사용한 claim 값(SOURCE_EVIDENCE_MATRIX.md 기록)이 DOCX 텍스트와 PDF 텍스트에 모두 나타나는지
2. DOCX 의 제목(장·절) 순서가 PDF 텍스트에 같은 순서로 나타나는지
3. 그림·표 캡션 번호가 연속인지, PDF 글꼴이 모두 임베드되었는지
결과: ``analysis/consistency_check.json`` (실패 항목 포함), 실패가 있으면 종료 코드 1.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pymupdf
from docx import Document

HERE = Path(__file__).resolve().parents[1]
DOCX = HERE / "R2H_T11_SVR_GARCH_CUSUM_최종연구보고서.docx"
PDF = HERE / "R2H_T11_SVR_GARCH_CUSUM_최종연구보고서.pdf"


def norm(s: str) -> str:
    return re.sub(r"\s+", "", s)


def docx_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main() -> int:
    claims = json.loads((HERE / "claims.json").read_text())
    used = re.findall(r"^\| `([^`]+)` \|", (HERE / "SOURCE_EVIDENCE_MATRIX.md").read_text(), flags=re.M)
    doc = Document(DOCX)
    dtext = norm(docx_text(doc))
    pdf = pymupdf.open(PDF)
    ptext = norm("".join(p.get_text() for p in pdf))
    out: dict[str, object] = {"pages": pdf.page_count, "claims_checked": 0, "missing_in_docx": [], "missing_in_pdf": []}
    for k in used:
        val = claims[k]["value"]
        if isinstance(val, dict):
            cands = [str(x) for x in val.values() if x is not None]
        elif isinstance(val, float):
            txt = f"{val:.6g}"
            cands = ([txt if len(txt) <= 7 else f"{val:.4g}"] if abs(val) >= 1e-3 or val == 0 else ["10⁻⁵미만"])
        else:
            cands = [str(val)]
        for c in cands:
            c = norm(c)
            out["claims_checked"] += 1
            if c not in dtext:
                out["missing_in_docx"].append(f"{k}={c}")
            if c not in ptext:
                out["missing_in_pdf"].append(f"{k}={c}")
    heads = [p.text for p in doc.paragraphs if p.style.name in ("Heading 1", "Heading 2") and p.text.strip()]
    pos, order_ok, missing_heads = 0, True, []
    body = ptext[ptext.find(norm("요약 (Executive Summary)"), 2000):]  # 목차 뒤부터
    for h in heads:
        i = body.find(norm(h), pos)
        if i < 0:
            missing_heads.append(h)
            order_ok = False
        else:
            pos = i
    out["headings"] = len(heads)
    out["heading_order_ok"] = order_ok
    out["headings_missing_in_pdf"] = missing_heads
    figs = [int(x) for x in re.findall(r"그림(\d+)\.", ptext)]
    tabs = [int(x) for x in re.findall(r"표(\d+)\.", ptext)]
    out["figure_captions"] = sorted(set(figs))
    out["table_captions"] = sorted(set(tabs))
    out["figures_contiguous"] = sorted(set(figs)) == list(range(1, max(figs) + 1))
    out["tables_contiguous"] = sorted(set(tabs)) == list(range(1, max(tabs) + 1))
    fonts = {f[3]: f[1] for p in pdf for f in p.get_fonts()}
    out["fonts"] = sorted(fonts)
    out["all_fonts_embedded"] = all(ext in ("ttf", "cff", "otf", "type1c", "cid") for ext in fonts.values()) and all(
        "+" in name for name in fonts)
    ok = (not out["missing_in_docx"] and not out["missing_in_pdf"] and order_ok and out["figures_contiguous"]
          and out["tables_contiguous"] and out["all_fonts_embedded"])
    out["PASS"] = ok
    (HERE / "analysis" / "consistency_check.json").write_text(json.dumps(out, ensure_ascii=False, indent=1))
    print(json.dumps(out, ensure_ascii=False)[:3000])
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
