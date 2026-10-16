"""T12 보고서 빌드 — 템플릿(Markdown + 수치 토큰) → 최종 Markdown · HTML · PDF(headless Chrome) · DOCX(python-docx).

수치는 원고에 직접 쓰지 않는다. ``{{표|키=값,…|열|형식}}`` 토큰을 결과표에서 찾아 치환하고, 치환할 때마다
SOURCE_EVIDENCE_MATRIX 에 (주장 ID, 위치, 표, 키, 열, 원값, 표시값) 을 남긴다 (명세 §8.1 #11, §13 #19).
실행: 상위 hongik ``.venv`` (python-docx · pymupdf 필요; R2H uv.lock 은 건드리지 않는다).

    python build_report.py --results ../results --figs ../figures --out OUTDIR
"""

from __future__ import annotations

import argparse
import html
import re
import subprocess
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
TOKEN = re.compile(r"\{\{([a-z_0-9]+)\|([^|}]*)\|([A-Za-z_0-9]+)\|([^}]*)\}\}")
ALIAS = {
    "vol": "annual_vol_diagnostics",
    "fc": "annual_forecast_diagnostics",
    "res": "annual_residual_diagnostics",
    "ph": "annual_posthoc",
    "cu": "annual_cusum_diagnostics",
    "ab": "annual_history_ablation",
    "pr": "diagnostic_profile",
    "fit": "fit_summary",
    "sy": "summary_by_year_policy",
    "syn": "null_dgp_validation",
    "par": "parallel_reproducibility",
    "agg": "aggregates",
}
YEARS = list(range(2020, 2027))


class Renderer:
    def __init__(self, results: Path):
        self.t = {}
        for k, v in ALIAS.items():
            p = results / f"{v}.csv"
            if p.exists():
                self.t[k] = pd.read_csv(p, dtype={"ticker": str})
        self.evidence: list[dict] = []
        self.section = ""

    def value(self, tbl: str, keys: str, col: str):
        df = self.t[tbl]
        m = pd.Series(True, index=df.index)
        for kv in filter(None, keys.split(",")):
            k, v = kv.split("=", 1)
            col_ = df[k]
            if pd.api.types.is_numeric_dtype(col_):
                m &= col_ == float(v)
            else:
                m &= col_.astype(str) == v
        rows = df[m]
        if len(rows) != 1:
            raise KeyError(f"token {tbl}|{keys}|{col} matched {len(rows)} rows")
        return rows.iloc[0][col]

    def fmt(self, v, f: str) -> str:
        if f == "s":
            return str(v)
        if f == "yn":
            return "기각" if bool(v) else "비기각"
        if pd.isna(v):
            return "NOT_ESTIMABLE"
        if f.endswith("%"):
            return format(float(v), f[:-1] + "%")
        if f == "d":
            return str(int(v))
        return format(float(v), f)

    def sub(self, text: str) -> str:
        def rep(m: re.Match) -> str:
            tbl, keys, col, f = m.groups()
            v = self.value(tbl, keys, col)
            out = self.fmt(v, f)
            self.evidence.append(
                {
                    "claim_id": f"C{len(self.evidence) + 1:04d}",
                    "section": self.section,
                    "table": ALIAS[tbl] + ".csv",
                    "keys": keys,
                    "column": col,
                    "raw_value": v,
                    "rendered": out,
                }
            )
            return out

        return TOKEN.sub(rep, text)


# ── Markdown 부분집합 → HTML / DOCX ───────────────────────────────────────────────
INLINE = [(re.compile(r"\*\*(.+?)\*\*"), r"<b>\1</b>"), (re.compile(r"`([^`]+)`"), r"<code>\1</code>")]


def inline_html(s: str) -> str:
    s = html.escape(s, quote=False)
    for p, r in INLINE:
        s = p.sub(r, s)
    return s


def blocks(md: str):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue
        if ln.strip() == "<!-- pagebreak -->":
            yield ("pagebreak", None)
            i += 1
            continue
        if ln.startswith("#"):
            lvl = len(ln) - len(ln.lstrip("#"))
            yield ("h", (lvl, ln[lvl:].strip()))
            i += 1
            continue
        if ln.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", ln)
            yield ("img", (m.group(1), m.group(2)))
            i += 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
            yield ("table", rows)
            continue
        if ln.startswith("> "):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i][1:].strip())
                i += 1
            yield ("quote", " ".join(buf))
            continue
        if re.match(r"^(- |\d+\. )", ln):
            ordered = bool(re.match(r"^\d+\. ", ln))
            items = []
            while i < len(lines) and re.match(r"^(- |\d+\. |  )", lines[i]):
                if re.match(r"^(- |\d+\. )", lines[i]):
                    items.append(re.sub(r"^(- |\d+\. )", "", lines[i]))
                else:
                    items[-1] += " " + lines[i].strip()
                i += 1
            yield ("list", (ordered, items))
            continue
        if ln.strip() == "---":
            yield ("hr", None)
            i += 1
            continue
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|!\[|\||> |- |\d+\. |---|<!--)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        yield ("p", " ".join(buf))


CSS = """
@font-face { font-family: 'NG'; src: url('file:///usr/share/fonts/truetype/nanum/NanumGothic.ttf'); }
@font-face { font-family: 'NG'; font-weight: bold; src: url('file:///usr/share/fonts/truetype/nanum/NanumGothicBold.ttf'); }
@page { size: A4; margin: 18mm 16mm 18mm 16mm; @bottom-center { content: counter(page); font-size: 8pt; } }
body { font-family: 'NG', sans-serif; font-size: 9.6pt; line-height: 1.55; color: #0b0b0b; background: #fff; }
h1 { font-size: 18pt; margin: 0 0 6pt; } h2 { font-size: 13.5pt; border-bottom: 1.5px solid #2a78d6; padding-bottom: 2pt; margin-top: 16pt; }
h3 { font-size: 11pt; margin: 12pt 0 4pt; color: #184f95; } h4 { font-size: 10pt; margin: 9pt 0 3pt; }
p { margin: 4pt 0; text-align: justify; } li { margin: 1.5pt 0; }
table { border-collapse: collapse; width: 100%; font-size: 7.6pt; margin: 6pt 0; page-break-inside: auto; }
th, td { border: 0.5px solid #c3c2b7; padding: 2pt 3pt; text-align: center; } th { background: #eef4fc; }
td:first-child { text-align: left; } tr { page-break-inside: avoid; }
img { width: 100%; margin: 4pt 0; page-break-inside: avoid; } figure { margin: 6pt 0; page-break-inside: avoid; }
figcaption { font-size: 8pt; color: #52514e; }
blockquote { border-left: 3px solid #eb6834; margin: 6pt 0; padding: 2pt 8pt; background: #fbf4ef; }
code { font-size: 8.4pt; background: #f0efec; padding: 0 2px; }
.pb { page-break-after: always; }
"""


def to_html(md: str, base: Path, title: str) -> str:
    out = [f"<!doctype html><html lang='ko'><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>{CSS}</style></head><body>"]
    for kind, v in blocks(md):
        if kind == "h":
            out.append(f"<h{v[0]}>{inline_html(v[1])}</h{v[0]}>")
        elif kind == "p":
            out.append(f"<p>{inline_html(v)}</p>")
        elif kind == "img":
            out.append(f"<figure><img src='{(base / v[1]).resolve().as_uri()}'><figcaption>{inline_html(v[0])}</figcaption></figure>")
        elif kind == "table":
            head, *body = v
            out.append("<table><thead><tr>" + "".join(f"<th>{inline_html(c)}</th>" for c in head) + "</tr></thead><tbody>")
            out += ["<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in body]
            out.append("</tbody></table>")
        elif kind == "quote":
            out.append(f"<blockquote>{inline_html(v)}</blockquote>")
        elif kind == "list":
            tag = "ol" if v[0] else "ul"
            out.append(f"<{tag}>" + "".join(f"<li>{inline_html(x)}</li>" for x in v[1]) + f"</{tag}>")
        elif kind == "hr":
            out.append("<hr>")
        elif kind == "pagebreak":
            out.append("<div class='pb'></div>")
    out.append("</body></html>")
    return "\n".join(out)


def to_docx(md: str, base: Path, path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    def font(style, size=None, bold=None, color=None):
        style.font.name = "NanumGothic"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "NanumGothic")
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color:
            style.font.color.rgb = RGBColor.from_string(color)

    font(doc.styles["Normal"], 9.5)
    for lvl, sz in ((1, 17), (2, 13), (3, 11), (4, 10)):
        font(doc.styles[f"Heading {lvl}"], sz, True, "184F95" if lvl > 1 else "0B0B0B")

    def runs(par, text):
        for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = "DejaVu Sans Mono"
                r.font.size = Pt(8)
            else:
                par.add_run(part)

    for kind, v in blocks(md):
        if kind == "h":
            doc.add_heading(re.sub(r"\*\*|`", "", v[1]), level=min(v[0], 4))
        elif kind == "p":
            runs(doc.add_paragraph(), v)
        elif kind == "img":
            doc.add_picture(str(base / v[1]), width=Cm(17.2))
            cap = doc.add_paragraph()
            r = cap.add_run(re.sub(r"\*\*|`", "", v[0]))
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string("52514E")
        elif kind == "table":
            head, *body = v
            tb = doc.add_table(rows=1, cols=len(head))
            tb.style = "Table Grid"
            for j, c in enumerate(head):
                cell = tb.rows[0].cells[j]
                cell.text = ""
                rr = cell.paragraphs[0].add_run(re.sub(r"\*\*|`", "", c))
                rr.bold = True
                rr.font.size = Pt(7)
            for row in body:
                cells = tb.add_row().cells
                for j, c in enumerate(row[: len(head)]):
                    cells[j].text = ""
                    rr = cells[j].paragraphs[0].add_run(re.sub(r"\*\*|`", "", c))
                    rr.font.size = Pt(7)
        elif kind == "quote":
            p = doc.add_paragraph(style="Intense Quote")
            runs(p, v)
        elif kind == "list":
            for x in v[1]:
                runs(doc.add_paragraph(style="List Number" if v[0] else "List Bullet"), x)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.core_properties.title = "T12 20ETF Annual Expanding Full Research Report"
    doc.core_properties.author = "R2H 2조 (T12 연구 에이전트 산출)"
    doc.save(path)


def render_pdf(html_path: Path, pdf_path: Path) -> None:
    subprocess.run(
        [
            "google-chrome",
            "--headless=new",
            "--disable-gpu",
            "--no-sandbox",
            "--allow-file-access-from-files",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path.as_uri(),
        ],
        check=True,
        capture_output=True,
    )


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", required=True)
    ap.add_argument("--figs", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", default=str(HERE / "REPORT_TEMPLATE.md"))
    ap.add_argument("--profiles", default=str(HERE / "etf_profiles.md"))
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    R = Renderer(Path(a.results))
    from sections import expand_directives  # noqa: E402  (같은 폴더)

    tmpl = Path(a.template).read_text(encoding="utf-8")
    md = expand_directives(tmpl, R, Path(a.results), Path(a.figs), Path(a.profiles))
    final = []
    for chunk in re.split(r"(?m)^(?=#{1,3} )", md):
        m = re.match(r"#{1,3} (.*)", chunk)
        if m:
            R.section = m.group(1)[:60]
        final.append(R.sub(chunk))
    md = "".join(final)
    left = re.findall(r"\{\{[^}]*\}\}", md)
    if left:
        raise SystemExit(f"unresolved tokens: {left[:5]}")
    name = "T12_20ETF_Annual_Expanding_Full_Research_Report"
    (out / f"{name}.md").write_text(md, encoding="utf-8")
    base = Path(a.figs).parent
    (out / f"{name}.html").write_text(to_html(md, base, name), encoding="utf-8")
    render_pdf(out / f"{name}.html", out / f"{name}.pdf")
    to_docx(md, base, out / f"{name}.docx")
    ev = pd.DataFrame(R.evidence)
    ev.to_csv(out / "SOURCE_EVIDENCE_MATRIX.csv", index=False, lineterminator="\n")
    print(f"tokens resolved: {len(ev)} → {out}")
